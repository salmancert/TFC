"""Reconcile a workbook of reconciliation packs, one sheet per bank.

Each sheet is self-contained: the open items from the accounting system and
the open items from the bank statement, stacked in one sheet, to be matched
against each other.  Sheets are processed independently and the results are
written back onto a copy of the original workbook, so the provider's layout,
formulas and formatting all survive and only the marking is added.
"""

from __future__ import annotations

import logging
import os
import shutil
from dataclasses import dataclass, field

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .highlight import (
    FILL_GROUPED,
    FILL_HEADER,
    FILL_MATCHED,
    FILL_REVIEW,
    FILL_UNMATCHED,
    FONT_GROUPED,
    FONT_HEADER,
    FONT_MATCHED,
    FONT_REVIEW,
    FONT_UNMATCHED,
    LEGEND,
)
from .matcher import ReconciliationConfig, ReconciliationResult, reconcile
from .single_sheet import SheetLayout, detect_layout, read_grid

LOGGER = logging.getLogger(__name__)

MARK_COLUMNS = ("Match ID", "Match Status", "Confidence", "Matched With")

_STYLES = {
    "matched": (FILL_MATCHED, FONT_MATCHED),
    "grouped": (FILL_GROUPED, FONT_GROUPED),
    "review": (FILL_REVIEW, FONT_REVIEW),
    "unmatched": (FILL_UNMATCHED, FONT_UNMATCHED),
}


@dataclass
class SheetOutcome:
    """What happened to one sheet of the pack."""

    sheet: str
    layout: SheetLayout | None = None
    result: ReconciliationResult | None = None
    skipped: str = ""

    @property
    def ok(self) -> bool:
        return self.result is not None

    @property
    def counts(self) -> dict[str, int]:
        if self.result is None:
            return {}
        records = self.result.records
        return {
            "matched": sum(1 for r in records if r.status == "matched" and r.kind == "one_to_one"),
            "grouped": sum(1 for r in records if r.status == "matched" and r.kind == "grouped"),
            "review": sum(1 for r in records if r.status == "review"),
            "unmatched_ledger": len(self.result.unmatched_left),
            "unmatched_bank": len(self.result.unmatched_right),
        }

    def describe(self) -> str:
        if not self.ok:
            return f"{self.sheet}: skipped - {self.skipped}"
        counts = self.counts
        return (
            f"{self.sheet}: {counts['matched'] + counts['grouped']} matched, "
            f"{counts['review']} to review, "
            f"{counts['unmatched_ledger'] + counts['unmatched_bank']} still open"
        )


@dataclass
class PackOutcome:
    """The whole workbook."""

    source_path: str
    sheets: list[SheetOutcome] = field(default_factory=list)
    output_path: str = ""
    notes: list[str] = field(default_factory=list)

    @property
    def processed(self) -> list[SheetOutcome]:
        return [outcome for outcome in self.sheets if outcome.ok]

    @property
    def totals(self) -> dict[str, int]:
        totals = {
            "matched": 0, "grouped": 0, "review": 0,
            "unmatched_ledger": 0, "unmatched_bank": 0,
        }
        for outcome in self.processed:
            for key, value in outcome.counts.items():
                totals[key] += value
        return totals

    def summary_frame(self) -> pd.DataFrame:
        rows = []
        for outcome in self.sheets:
            if not outcome.ok:
                rows.append({"Sheet": outcome.sheet, "Status": f"skipped - {outcome.skipped}"})
                continue
            counts = outcome.counts
            assert outcome.result is not None
            rows.append({
                "Sheet": outcome.sheet,
                "Status": "reconciled",
                "Matched": counts["matched"],
                "Grouped": counts["grouped"],
                "To review": counts["review"],
                f"Still open ({outcome.result.left.name})": counts["unmatched_ledger"],
                f"Still open ({outcome.result.right.name})": counts["unmatched_bank"],
                "Scoring model": outcome.result.training.summary(),
            })
        return pd.DataFrame(rows)


def sheet_names(path: str) -> list[str]:
    """Every sheet in the workbook, in order."""
    if str(path).lower().endswith((".csv", ".txt", ".tsv")):
        return [os.path.basename(path)]
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def reconcile_sheet(
    grid: pd.DataFrame, sheet: str, config: ReconciliationConfig | None = None
) -> SheetOutcome:
    """Find the two blocks in one sheet and match them against each other."""
    layout = detect_layout(grid)
    outcome = SheetOutcome(sheet=sheet, layout=layout)
    if not layout.is_pack:
        outcome.skipped = (
            layout.notes[-1] if layout.notes else "no two blocks of open items found"
        )
        return outcome

    ledger, bank = layout.sections
    outcome.result = reconcile(
        ledger.frame, bank.frame, ledger.label, bank.label, config=config
    )
    return outcome


def reconcile_pack(
    path: str,
    config: ReconciliationConfig | None = None,
    sheets: list[str] | None = None,
) -> PackOutcome:
    """Reconcile every sheet of a workbook independently."""
    outcome = PackOutcome(source_path=path)
    targets = sheets if sheets is not None else sheet_names(path)
    for sheet in targets:
        try:
            grid = read_grid(path, sheet)
        except Exception as error:  # noqa: BLE001 - reported per sheet
            outcome.sheets.append(SheetOutcome(sheet=sheet, skipped=f"could not read: {error}"))
            continue
        if grid.empty:
            outcome.sheets.append(SheetOutcome(sheet=sheet, skipped="sheet is empty"))
            continue
        try:
            outcome.sheets.append(reconcile_sheet(grid, sheet, config))
        except Exception as error:  # noqa: BLE001 - one bad sheet must not stop the rest
            LOGGER.exception("Sheet %r failed", sheet)
            outcome.sheets.append(SheetOutcome(sheet=sheet, skipped=f"failed: {error}"))
    return outcome


def _row_styles(result: ReconciliationResult) -> tuple[dict[int, tuple], dict[int, tuple]]:
    """Per section-row: (style key, match id, confidence, partner rows)."""
    left: dict[int, tuple] = {}
    right: dict[int, tuple] = {}
    for number, record in enumerate(result.records, start=1):
        match_id = f"M{number:04d}"
        style = (
            "review" if record.status == "review"
            else "grouped" if record.kind == "grouped"
            else "matched"
        )
        for row in record.left_rows:
            left[row] = (style, match_id, record.score, record.right_rows)
        for row in record.right_rows:
            right[row] = (style, match_id, record.score, record.left_rows)
    for row in result.unmatched_left:
        left[row] = ("unmatched", "", None, ())
    for row in result.unmatched_right:
        right[row] = ("unmatched", "", None, ())
    return left, right


def mark_workbook(outcome: PackOutcome, output_path: str) -> str:
    """Copy the pack and write the marking onto it, layout untouched."""
    import openpyxl

    source = outcome.source_path
    if os.path.abspath(source) == os.path.abspath(output_path):
        raise ValueError("The marked copy would overwrite the source pack.")
    directory = os.path.dirname(os.path.abspath(output_path))
    if directory:
        os.makedirs(directory, exist_ok=True)
    shutil.copyfile(source, output_path)

    book = openpyxl.load_workbook(output_path)
    for sheet_outcome in outcome.processed:
        _mark_sheet(book[sheet_outcome.sheet], sheet_outcome)
    _add_summary_sheet(book, outcome)
    book.save(output_path)
    outcome.output_path = output_path
    return output_path


def _mark_sheet(worksheet, sheet_outcome: SheetOutcome) -> None:
    """Colour the item rows of one sheet and append the marking columns."""
    result = sheet_outcome.result
    layout = sheet_outcome.layout
    assert result is not None and layout is not None
    ledger_section, bank_section = layout.sections
    left_styles, right_styles = _row_styles(result)

    # Put the marking columns clear of whatever the provider already uses.
    first_mark_column = worksheet.max_column + 2
    width = first_mark_column + len(MARK_COLUMNS) - 1

    for section, styles, partner in (
        (ledger_section, left_styles, bank_section),
        (bank_section, right_styles, ledger_section),
    ):
        if section.header_row is not None:
            header_excel_row = section.header_row + 1
            for offset, caption in enumerate(MARK_COLUMNS):
                cell = worksheet.cell(
                    row=header_excel_row, column=first_mark_column + offset, value=caption
                )
                cell.fill = FILL_HEADER
                cell.font = FONT_HEADER

        for position, source_row in enumerate(section.source_rows):
            style, match_id, confidence, partner_rows = styles.get(
                position, ("unmatched", "", None, ())
            )
            fill, font = _STYLES[style]
            excel_row = source_row + 1

            for column in range(1, width + 1):
                cell = worksheet.cell(row=excel_row, column=column)
                cell.fill = fill
                if cell.value is not None:
                    cell.font = Font(
                        color=font.color.rgb if font.color else None,
                        bold=cell.font.bold,
                        italic=cell.font.italic,
                        size=cell.font.size,
                        name=cell.font.name,
                    )

            partner_label = (
                ", ".join(str(partner.source_rows[r] + 1) for r in partner_rows)
                if partner_rows else ""
            )
            values = (
                match_id,
                style,
                round(float(confidence), 3) if confidence is not None else None,
                f"row {partner_label}" if partner_label else "",
            )
            for offset, value in enumerate(values):
                worksheet.cell(row=excel_row, column=first_mark_column + offset, value=value)

    for offset in range(len(MARK_COLUMNS)):
        letter = get_column_letter(first_mark_column + offset)
        worksheet.column_dimensions[letter].width = 16 if offset < 3 else 22


def _add_summary_sheet(book, outcome: PackOutcome, title: str = "Reconciliation Summary") -> None:
    """One overview sheet across every bank in the pack."""
    if title in book.sheetnames:
        del book[title]
    sheet = book.create_sheet(title, 0)

    sheet.cell(row=1, column=1, value="Reconciliation Summary").font = Font(bold=True, size=14)
    sheet.cell(row=2, column=1, value=f"Source: {os.path.basename(outcome.source_path)}")

    row = 4
    for style, meaning, fill, font in LEGEND:
        sheet.cell(row=row, column=1, value=style).fill = fill
        sheet.cell(row=row, column=1).font = font
        sheet.cell(row=row, column=2, value=meaning)
        row += 1

    row += 1
    frame = outcome.summary_frame()
    if not frame.empty:
        for offset, column in enumerate(frame.columns, start=1):
            cell = sheet.cell(row=row, column=offset, value=str(column))
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
        for _, record in frame.iterrows():
            row += 1
            for offset, column in enumerate(frame.columns, start=1):
                value = record[column]
                sheet.cell(
                    row=row, column=offset,
                    value=None if pd.isna(value) else (
                        value if isinstance(value, (int, float, str)) else str(value)
                    ),
                )

    row += 2
    totals = outcome.totals
    sheet.cell(row=row, column=1, value="Totals across all sheets").font = Font(bold=True)
    for label, key in (
        ("Matched", "matched"), ("Grouped", "grouped"), ("To review", "review"),
        ("Still open (ledger side)", "unmatched_ledger"),
        ("Still open (bank side)", "unmatched_bank"),
    ):
        row += 1
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=totals[key])

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 46
    for offset in range(3, 10):
        sheet.column_dimensions[get_column_letter(offset)].width = 18
