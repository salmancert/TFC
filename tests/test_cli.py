"""The command line interface."""

import openpyxl
import pandas as pd
import pytest

from bank_reconciliation.cli import main
from bank_reconciliation.sample_data import write_sample_workbook


@pytest.fixture
def workbook(tmp_path):
    return write_sample_workbook(str(tmp_path / "statement.xlsx"), n_matched=30, seed=17)


def test_default_run_writes_a_coloured_copy(workbook, tmp_path, capsys):
    assert main([workbook]) == 0
    output = tmp_path / "statement_reconciled.xlsx"
    assert output.exists()
    book = openpyxl.load_workbook(output)
    assert "Legend" in book.sheetnames
    assert "Auto-matched pairs" in capsys.readouterr().out


def test_source_file_is_left_untouched(workbook):
    before = openpyxl.load_workbook(workbook).sheetnames
    main([workbook])
    assert openpyxl.load_workbook(workbook).sheetnames == before


def test_tables_mode_writes_the_plain_report(workbook, tmp_path):
    target = tmp_path / "tables.xlsx"
    assert main([workbook, "--tables", "-o", str(target)]) == 0
    assert set(openpyxl.load_workbook(target).sheetnames) >= {"Summary", "Matches"}


def test_in_place_appends_to_the_source_workbook(workbook):
    assert main([workbook, "--in-place"]) == 0
    names = openpyxl.load_workbook(workbook).sheetnames
    assert "Bank_Statement" in names and "Summary" in names


def test_refuses_to_overwrite_the_input(workbook):
    assert main([workbook, "-o", workbook]) == 2


def test_missing_file_is_reported(tmp_path):
    assert main([str(tmp_path / "nope.xlsx")]) == 2


def test_contradictory_thresholds_are_rejected(workbook):
    assert main([workbook, "--match-threshold", "0.2", "--review-threshold", "0.9"]) == 2


def test_two_separate_files(tmp_path):
    from bank_reconciliation.sample_data import generate_sample

    sample = generate_sample(n_matched=25, seed=19)
    bank = tmp_path / "bank.csv"
    ledger = tmp_path / "ledger.csv"
    sample.bank.to_csv(bank, index=False)
    sample.ledger.to_csv(ledger, index=False)
    target = tmp_path / "out.xlsx"
    assert main([str(bank), "--ledger", str(ledger), "-o", str(target)]) == 0
    assert target.exists()


def test_csv_dir_output(workbook, tmp_path):
    csv_dir = tmp_path / "csvs"
    assert main([workbook, "--csv-dir", str(csv_dir)]) == 0
    assert len(list(csv_dir.glob("*.csv"))) == 4


def test_explicit_sheets(workbook, tmp_path):
    target = tmp_path / "explicit.xlsx"
    assert main([workbook, "--sheets", "Internal_Ledger", "Bank_Statement",
                 "-o", str(target)]) == 0
    assert target.exists()


def test_unknown_sheet_is_reported(workbook, tmp_path):
    with pytest.raises(SystemExit):
        main([workbook, "--sheets", "Nope", "Bank_Statement"])


def test_model_can_be_saved_and_reused(workbook, tmp_path):
    model = tmp_path / "scorer.joblib"
    assert main([workbook, "--save-model", str(model), "-o", str(tmp_path / "a.xlsx")]) == 0
    if model.exists():  # only written when a model could actually be trained
        assert main([workbook, "--load-model", str(model),
                     "-o", str(tmp_path / "b.xlsx")]) == 0


def test_single_sheet_workbook_is_reported(tmp_path):
    target = tmp_path / "one.xlsx"
    pd.DataFrame({"Date": ["2024-01-02"], "Memo": ["x"], "Amount": [-1.0]}).to_excel(
        target, sheet_name="Only", index=False
    )
    with pytest.raises(SystemExit):
        main([str(target)])
