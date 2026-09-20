"""Reconciling a workbook of packs, one sheet per bank, and marking it."""

import openpyxl
import pandas as pd
import pytest

from bank_reconciliation.highlight import FILL_MATCHED, FILL_UNMATCHED
from bank_reconciliation.pack import (
    MARK_COLUMNS,
    mark_workbook,
    reconcile_pack,
    reconcile_sheet,
    sheet_names,
)
from bank_reconciliation.sample_data import generate_pack, write_sample_multi_bank_pack


@pytest.fixture
def multi_bank(tmp_path):
    path, truths = write_sample_multi_bank_pack(str(tmp_path / "banks.xlsx"))
    return path, truths


def test_every_sheet_is_reconciled_independently(multi_bank):
    path, truths = multi_bank
    outcome = reconcile_pack(path)
    assert len(outcome.processed) == len(truths) == 3
    for sheet_outcome in outcome.processed:
        assert sheet_outcome.result is not None
        assert sheet_outcome.counts["matched"] > 0


def test_matches_are_correct_against_known_answers(multi_bank):
    path, truths = multi_bank
    outcome = reconcile_pack(path)
    for sheet_outcome in outcome.processed:
        truth = truths[sheet_outcome.sheet]
        ledger, bank = sheet_outcome.layout.sections
        predicted = {
            ledger.source_rows[record.left_rows[0]]: bank.source_rows[record.right_rows[0]]
            for record in sheet_outcome.result.records
            if len(record.left_rows) == 1 and len(record.right_rows) == 1
        }
        correct = sum(1 for k, v in predicted.items() if truth.get(k) == v)
        assert correct / max(len(predicted), 1) >= 0.98
        assert correct / len(truth) >= 0.95


def test_totals_add_up_across_sheets(multi_bank):
    path, _ = multi_bank
    outcome = reconcile_pack(path)
    totals = outcome.totals
    per_sheet = sum(
        o.counts["matched"] + o.counts["grouped"] for o in outcome.processed
    )
    assert totals["matched"] + totals["grouped"] == per_sheet


def test_only_selected_sheets_are_processed(multi_bank):
    path, _ = multi_bank
    chosen = sheet_names(path)[1]
    outcome = reconcile_pack(path, sheets=[chosen])
    assert [o.sheet for o in outcome.sheets] == [chosen]


def test_a_sheet_that_is_not_a_pack_is_skipped_without_stopping_the_rest(tmp_path):
    target = tmp_path / "mixed.xlsx"
    sheet, _ = generate_pack()
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        sheet.to_excel(writer, sheet_name="Good", index=False, header=False)
        pd.DataFrame({"Notes": ["nothing to reconcile here"]}).to_excel(
            writer, sheet_name="Cover", index=False
        )
    outcome = reconcile_pack(str(target))
    assert [o.sheet for o in outcome.processed] == ["Good"]
    skipped = [o for o in outcome.sheets if not o.ok]
    assert len(skipped) == 1 and skipped[0].skipped


# ------------------------------------------------------------------- marking
def test_marking_leaves_the_source_workbook_untouched(multi_bank, tmp_path):
    path, _ = multi_bank
    before = openpyxl.load_workbook(path)
    before_rows = {
        name: [[c.value for c in row] for row in before[name].iter_rows()]
        for name in before.sheetnames
    }
    outcome = reconcile_pack(path)
    mark_workbook(outcome, str(tmp_path / "marked.xlsx"))

    after = openpyxl.load_workbook(path)
    assert after.sheetnames == before.sheetnames
    for name in after.sheetnames:
        assert [[c.value for c in row] for row in after[name].iter_rows()] == before_rows[name]


def test_marking_refuses_to_overwrite_the_source(multi_bank):
    path, _ = multi_bank
    outcome = reconcile_pack(path)
    with pytest.raises(ValueError, match="overwrite"):
        mark_workbook(outcome, path)


def test_original_cells_survive_marking(multi_bank, tmp_path):
    path, _ = multi_bank
    outcome = reconcile_pack(path)
    target = str(tmp_path / "marked.xlsx")
    mark_workbook(outcome, target)

    source = openpyxl.load_workbook(path)["HSBC Current"]
    marked = openpyxl.load_workbook(target)["HSBC Current"]
    for row in range(1, source.max_row + 1):
        for column in range(1, source.max_column + 1):
            assert marked.cell(row=row, column=column).value == \
                source.cell(row=row, column=column).value


def test_item_rows_are_coloured_and_other_rows_are_not(multi_bank, tmp_path):
    path, _ = multi_bank
    outcome = reconcile_pack(path)
    target = str(tmp_path / "marked.xlsx")
    mark_workbook(outcome, target)

    sheet_outcome = outcome.processed[0]
    marked = openpyxl.load_workbook(target)[sheet_outcome.sheet]
    item_rows = {
        row + 1
        for section in sheet_outcome.layout.sections
        for row in section.source_rows
    }
    coloured = {
        row for row in range(1, marked.max_row + 1)
        if marked.cell(row=row, column=1).fill.start_color.rgb
        in {FILL_MATCHED.start_color.rgb, FILL_UNMATCHED.start_color.rgb}
        or marked.cell(row=row, column=1).fill.patternType == "solid"
    }
    # The balance lines, titles and header rows must stay unpainted.
    assert coloured <= item_rows


def test_marking_columns_are_added_clear_of_the_existing_data(multi_bank, tmp_path):
    path, _ = multi_bank
    outcome = reconcile_pack(path)
    target = str(tmp_path / "marked.xlsx")
    mark_workbook(outcome, target)

    sheet_outcome = outcome.processed[0]
    source = openpyxl.load_workbook(path)[sheet_outcome.sheet]
    marked = openpyxl.load_workbook(target)[sheet_outcome.sheet]
    first_mark = source.max_column + 2
    header_row = sheet_outcome.layout.sections[0].header_row + 1
    captions = [
        marked.cell(row=header_row, column=first_mark + offset).value
        for offset in range(len(MARK_COLUMNS))
    ]
    assert captions == list(MARK_COLUMNS)


def test_match_ids_are_reciprocal_between_the_two_blocks(multi_bank, tmp_path):
    path, _ = multi_bank
    outcome = reconcile_pack(path)
    target = str(tmp_path / "marked.xlsx")
    mark_workbook(outcome, target)

    sheet_outcome = outcome.processed[0]
    source = openpyxl.load_workbook(path)[sheet_outcome.sheet]
    marked = openpyxl.load_workbook(target)[sheet_outcome.sheet]
    id_column = source.max_column + 2
    partner_column = id_column + 3

    item_rows = {
        row + 1
        for section in sheet_outcome.layout.sections
        for row in section.source_rows
    }
    checked = 0
    for row in sorted(item_rows):
        match_id = marked.cell(row=row, column=id_column).value
        partner = marked.cell(row=row, column=partner_column).value
        if not match_id or not partner:
            continue
        for partner_row in [int(p) for p in str(partner).replace("row", "").split(",")]:
            assert marked.cell(row=partner_row, column=id_column).value == match_id
            assert str(row) in str(marked.cell(row=partner_row, column=partner_column).value)
            checked += 1
    assert checked > 0


def test_summary_sheet_lists_every_bank(multi_bank, tmp_path):
    path, truths = multi_bank
    outcome = reconcile_pack(path)
    target = str(tmp_path / "marked.xlsx")
    mark_workbook(outcome, target)

    book = openpyxl.load_workbook(target)
    assert book.sheetnames[0] == "Reconciliation Summary"
    summary = book["Reconciliation Summary"]
    text = " ".join(
        str(summary.cell(row=r, column=c).value)
        for r in range(1, summary.max_row + 1)
        for c in range(1, summary.max_column + 1)
    )
    for sheet in truths:
        assert sheet in text
    assert "Totals across all sheets" in text


def test_reconcile_sheet_works_on_a_bare_frame():
    sheet, truth = generate_pack()
    outcome = reconcile_sheet(sheet, "inline")
    assert outcome.ok
    assert outcome.counts["matched"] >= len(truth) * 0.95


def test_formulas_styling_and_widths_survive_marking(tmp_path):
    """The provider's own pack must come back intact apart from the marking."""
    from openpyxl.styles import Font

    target = tmp_path / "styled.xlsx"
    sheet, _ = generate_pack()
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        sheet.to_excel(writer, sheet_name="Bank A", index=False, header=False)

    book = openpyxl.load_workbook(target)
    worksheet = book["Bank A"]
    worksheet["A70"] = "Total SAP items"
    worksheet["D70"] = "=SUM(D9:D34)"
    worksheet["A1"].font = Font(bold=True, size=16)
    worksheet.column_dimensions["C"].width = 55
    book.save(target)

    outcome = reconcile_pack(str(target))
    marked_path = str(tmp_path / "styled_marked.xlsx")
    mark_workbook(outcome, marked_path)

    marked = openpyxl.load_workbook(marked_path)["Bank A"]
    assert marked["D70"].value == "=SUM(D9:D34)"
    assert marked["A1"].font.bold and marked["A1"].font.size == 16
    assert marked.column_dimensions["C"].width == 55
    # A total line is not an open item and must not be coloured.
    assert marked["A70"].fill.patternType is None
