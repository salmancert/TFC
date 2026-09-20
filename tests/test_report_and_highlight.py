"""Report tables and the colour-coded workbook."""

import openpyxl
import pandas as pd
import pytest

from bank_reconciliation import ReconciliationConfig, build_frames, reconcile, write_excel
from bank_reconciliation.highlight import (
    FILL_GROUPED,
    FILL_MATCHED,
    FILL_REVIEW,
    FILL_UNMATCHED,
    write_highlighted_workbook,
)
from bank_reconciliation.report import format_console_report, write_csvs
from bank_reconciliation.sample_data import generate_sample


@pytest.fixture(scope="module")
def result():
    sample = generate_sample(n_matched=40, seed=13)
    return reconcile(sample.bank, sample.ledger, "Bank_Statement", "Internal_Ledger")


def test_report_frames_cover_every_row_exactly_once(result):
    frames = build_frames(result)
    assert set(frames) == {
        "Summary", "Matches",
        "Unmatched_Bank_Statement", "Unmatched_Internal_Ledger",
    }
    matched_left = sum(len(r.left_rows) for r in result.records)
    assert matched_left + len(frames["Unmatched_Bank_Statement"]) == len(result.left)
    matched_right = sum(len(r.right_rows) for r in result.records)
    assert matched_right + len(frames["Unmatched_Internal_Ledger"]) == len(result.right)


def test_matches_table_shows_both_sides_and_the_difference(result):
    matches = build_frames(result)["Matches"]
    assert "Amount_Difference" in matches.columns
    assert "Bank_Statement.Amount" in matches.columns
    assert "Internal_Ledger.Net Amount" in matches.columns
    # Every accepted match should balance to the cent.
    accepted = matches[matches["Status"] == "matched"]
    assert accepted["Amount_Difference"].abs().max() <= 0.01


def test_row_numbers_point_at_the_spreadsheet_row(result):
    """Row 0 of the frame is row 2 of the sheet, under the header."""
    matches = build_frames(result)["Matches"]
    first = result.records[0]
    assert str(first.left_rows[0] + 2) in str(matches.iloc[0]["Bank_Statement_Row"])


def test_write_excel_produces_every_sheet(tmp_path, result):
    target = tmp_path / "report.xlsx"
    write_excel(result, str(target))
    book = openpyxl.load_workbook(target)
    assert "Summary" in book.sheetnames and "Matches" in book.sheetnames


def test_write_csvs(tmp_path, result):
    written = write_csvs(result, str(tmp_path / "csvs"))
    assert len(written) == 4
    assert all(pd.read_csv(path) is not None for path in written)


def test_console_report_mentions_the_headline_numbers(result):
    text = format_console_report(result)
    assert "Reconciliation summary" in text
    assert "Auto-matched pairs" in text


# --------------------------------------------------------------- highlighting
def test_highlighted_workbook_keeps_the_source_sheets(tmp_path, result):
    target = tmp_path / "coloured.xlsx"
    write_highlighted_workbook(result, str(target))
    book = openpyxl.load_workbook(target)
    assert book.sheetnames == ["Legend", "Bank_Statement", "Internal_Ledger", "Matches"]

    sheet = book["Bank_Statement"]
    # Original columns first, verdict columns appended.
    headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    assert headers[: len(result.left.frame.columns)] == list(result.left.frame.columns)
    assert headers[-4:] == ["Match ID", "Match Status", "Confidence", "Matched With"]
    assert sheet.max_row == len(result.left) + 1


def test_every_row_is_coloured_according_to_its_verdict(tmp_path, result):
    target = tmp_path / "coloured.xlsx"
    write_highlighted_workbook(result, str(target))
    sheet = openpyxl.load_workbook(target)["Bank_Statement"]

    expected = {}
    for record in result.records:
        fill = (
            FILL_REVIEW if record.status == "review"
            else FILL_GROUPED if record.kind == "grouped"
            else FILL_MATCHED
        )
        for row in record.left_rows:
            expected[row] = fill.start_color.rgb
    for row in result.unmatched_left:
        expected[row] = FILL_UNMATCHED.start_color.rgb

    for row_index, colour in expected.items():
        cell = sheet.cell(row=row_index + 2, column=1)
        assert cell.fill.start_color.rgb == colour, f"row {row_index + 2} has the wrong colour"


def test_matched_rows_carry_the_same_match_id_on_both_sides(tmp_path, result):
    target = tmp_path / "coloured.xlsx"
    write_highlighted_workbook(result, str(target))
    book = openpyxl.load_workbook(target)
    bank, ledger = book["Bank_Statement"], book["Internal_Ledger"]
    id_column = len(result.left.frame.columns) + 1

    record = result.records[0]
    bank_id = bank.cell(row=record.left_rows[0] + 2, column=id_column).value
    ledger_id_column = len(result.right.frame.columns) + 1
    ledger_id = ledger.cell(row=record.right_rows[0] + 2, column=ledger_id_column).value
    assert bank_id and bank_id == ledger_id


def test_unmatched_rows_have_no_match_id(tmp_path, result):
    target = tmp_path / "coloured.xlsx"
    write_highlighted_workbook(result, str(target))
    sheet = openpyxl.load_workbook(target)["Bank_Statement"]
    id_column = len(result.left.frame.columns) + 1
    for row in result.unmatched_left:
        assert not sheet.cell(row=row + 2, column=id_column).value


def test_legend_sheet_explains_the_colours(tmp_path, result):
    target = tmp_path / "coloured.xlsx"
    write_highlighted_workbook(result, str(target))
    sheet = openpyxl.load_workbook(target)["Legend"]
    text = " ".join(
        str(sheet.cell(row=r, column=c).value)
        for r in range(1, sheet.max_row + 1) for c in (1, 2)
    )
    assert "Matched automatically" in text and "Unmatched" in text
    assert "Auto-matched pairs" in text  # summary block


def test_sheet_names_with_illegal_characters_are_sanitised(tmp_path):
    bank = pd.DataFrame({"Date": ["2024-01-02"], "Memo": ["x"], "Amount": [-1.0]})
    ledger = pd.DataFrame({"Date": ["2024-01-02"], "Memo": ["x"], "Amount": [-1.0]})
    result = reconcile(bank, ledger, "Bank/Statement[2024]", "Ledger:main")
    target = tmp_path / "odd.xlsx"
    write_highlighted_workbook(result, str(target))
    names = openpyxl.load_workbook(target).sheetnames
    assert all(not set(name) & set(":*?/\\[]") for name in names)
    assert all(len(name) <= 31 for name in names)


def test_review_rows_are_amber(tmp_path):
    bank = pd.DataFrame({"Date": ["2024-01-02"], "Memo": ["UNKNOWN MERCHANT"], "Amount": [-100.0]})
    ledger = pd.DataFrame({"Date": ["2024-01-02"], "Memo": ["Something else"], "Amount": [-100.4]})
    config = ReconciliationConfig(match_threshold=0.95, review_threshold=0.2)
    result = reconcile(bank, ledger, "Bank", "Ledger", config=config)
    assert [r.status for r in result.records] == ["review"]
    target = tmp_path / "review.xlsx"
    write_highlighted_workbook(result, str(target))
    sheet = openpyxl.load_workbook(target)["Bank"]
    assert sheet.cell(row=2, column=1).fill.start_color.rgb == FILL_REVIEW.start_color.rgb
